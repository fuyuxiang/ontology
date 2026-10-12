"""SQLite persistence regressions for builder ownership and document mapping."""
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.api.v1 import builder, doc_builder, entities, ai_builder_v2, ontology_publish, ontology_api, scenarios, ontology_mapping, actions, ai_code, functions
from app.core.deps import require_user
from app.database import Base, get_db
from app.models.asset import Asset
from app.models.entity import OntologyEntity, EntityAttribute
from app.models.object_binding import ObjectBinding
from app.models.scenario import ScenarioDict
from app.models.user import User
from app.services import doc_builder as document_service


@pytest.fixture
def environment():
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine, expire_on_commit=False)
    with sessions() as db:
        db.add_all([ScenarioDict(id='oa', code='oa', name='A'), ScenarioDict(id='ob', code='ob', name='B')])
        db.add(Asset(id='asset', name='devices', kind='table', status='active', locator={'table': 'devices'}, schema_snapshot=[{'name': 'device_id', 'type': 'string'}, {'name': 'risk_score', 'type': 'number'}], primary_key=['device_id']))
        db.commit()
    app = FastAPI()
    app.include_router(builder.router)
    app.include_router(doc_builder.router)
    app.include_router(entities.router)
    app.include_router(ai_builder_v2.router)
    app.include_router(ontology_publish.router)
    app.include_router(ontology_api.router)
    app.include_router(scenarios.router)
    app.include_router(ontology_mapping.router)
    app.include_router(actions.router)
    app.include_router(ai_code.router)
    app.include_router(functions.router)
    def database():
        with sessions() as db:
            yield db
    app.dependency_overrides[get_db] = database
    app.dependency_overrides[require_user] = lambda: User(id='owner', username='owner', password_hash='x', name='Owner', role='admin', is_active=True)
    yield TestClient(app), sessions
    engine.dispose()


def draft(ontology_id, column='risk_score'):
    return {'ontology_id': ontology_id, 'objects': [{'name': 'Device', 'displayName': '设备', 'tier': 2, 'primaryKey': 'id', 'backing_asset_ids': ['asset'], 'properties': [{'name': 'id', 'type': 'string', 'required': True, 'source_asset_id': 'asset', 'source_column': 'device_id'}, {'name': 'risk', 'type': 'number', 'source_asset_id': 'asset', 'source_column': column}]}]}


def test_finalize_scopes_same_named_entities_and_updates_mapping(environment):
    client, sessions = environment
    a = client.post('/builder/finalize', json=draft('oa'))
    b = client.post('/builder/finalize', json=draft('ob'))
    assert a.status_code == b.status_code == 200
    assert a.json()['state'] == 'saved'
    assert a.json()['entity_ids']['Device'] != b.json()['entity_ids']['Device']
    assert client.post('/builder/finalize', json=draft('oa')).json()['created_entities'] == 0
    with sessions() as db:
        assert len(db.scalars(select(OntologyEntity)).all()) == 2
        bindings = db.scalars(select(ObjectBinding)).all()
        assert len(bindings) == 2
        assert all(binding.id_column == 'device_id' for binding in bindings)
        assert all(len(binding.field_mappings) == 2 for binding in bindings)


def test_finalize_invalid_asset_field_and_relation_do_not_write(environment):
    client, sessions = environment
    assert client.post('/builder/finalize', json=draft('oa', 'missing')).status_code == 422
    request = draft('oa')
    request['relations'] = [{'name': 'bad', 'source': 'Device', 'target': 'Unknown'}]
    assert client.post('/builder/finalize', json=request).status_code == 422
    with sessions() as db:
        assert db.scalar(select(OntologyEntity)) is None


def test_document_apply_creates_new_attributes_before_binding_and_isolates(environment):
    client, sessions = environment
    client.post('/builder/finalize', json=draft('ob')).raise_for_status()
    mapping = {'entities': [{'name': 'Device', 'displayName': '设备', 'asset_id': 'asset', 'table': 'devices', 'primaryKey': 'id', 'properties': [{'name': 'id', 'type': 'string', 'required': True, 'field': 'device_id'}, {'name': 'risk', 'type': 'number', 'field': 'risk_score'}]}], 'relations': []}
    preview = client.post('/doc-builder/mapping/preview', json={'session_id': 'doc', 'ontology_id': 'oa', 'mapping_result': mapping}).json()['items'][0]
    assert preview['entity_id'] is None
    result = client.post('/doc-builder/mapping/apply', json={'session_id': 'doc', 'ontology_id': 'oa', 'mapping_result': mapping, 'items': [{'entity_name': 'Device', 'entity_id': None, 'asset_id': 'asset', 'field_mappings': []}]})
    assert result.status_code == 200, result.text
    assert result.json()['created'] == 1 and result.json()['skipped'] == 0
    with sessions() as db:
        entity = db.scalar(select(OntologyEntity).where(OntologyEntity.ontology_id == 'oa'))
        attributes = db.scalars(select(EntityAttribute).where(EntityAttribute.entity_id == entity.id)).all()
        binding = db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id == entity.id))
        assert {field['attribute_id'] for field in binding.field_mappings} == {attribute.id for attribute in attributes}
        assert {attribute.source_field for attribute in attributes} == {'device_id', 'risk_score'}


def test_document_upload_survives_memory_reset_and_checks_owner(environment):
    client, sessions = environment
    result = client.post('/doc-builder/upload', files={'files': ('device.md', b'Device fields: id, risk', 'text/markdown')})
    assert result.status_code == 200
    session_id = result.json()['session_id']
    document_service._sessions.clear()
    assert client.get(f'/doc-builder/sessions/{session_id}').status_code == 200
    with sessions() as db:
        assert document_service.get_session(session_id, db, 'other') is None
        assert document_service.get_session(session_id, db, 'owner')['files'][0]['content'] == 'Device fields: id, risk'
    key = f'doc-upload:{session_id}:device.md'
    assert client.get('/ai-builder/documents/content', params={'key':key}).json()['content'] == 'Device fields: id, risk'
    from app.services.ai_builder_v2 import read_document_content
    with sessions() as db:
        with pytest.raises(ValueError,match='不可访问'):
            read_document_content(key,db,'other')


def test_editor_owl_export_reimports_core_definitions_and_rdf_xml_remains_compatible(environment):
    from pathlib import Path
    client, sessions = environment
    content = (Path(__file__).parent/'fixtures'/'builder-owl-roundtrip.owl').read_bytes()
    preview = client.post('/entities/preview-file',data={'file_type':'owl'},files={'file':('ontology.owl',content,'application/xml')})
    assert preview.status_code == 200,preview.text
    objects = {obj['name']:obj for obj in preview.json()['objects']}
    assert set(objects) == {'Asset','Device','WorkOrder'}
    assert objects['Device']['display_name'] == '设备 & 告警'
    assert objects['Device']['properties'][0]['name'] == 'risk'
    assert objects['Device']['properties'][0]['type'] == 'number'
    assert {(rel['name'],rel['cardinality']) for rel in preview.json()['relations']} == {('hasOrder','1:N'),('subClassOf','N:1')}
    for ontology_id in ['oa','ob']:
        saved = client.post('/entities/from-file',data={'file_type':'owl','ontology_id':ontology_id},files={'file':('ontology.owl',content,'application/xml')})
        assert saved.status_code == 201,saved.text
        assert saved.json()['entities_created'] == 3
        assert saved.json()['attributes_created'] == 1 and saved.json()['relations_created'] == 2
    with sessions() as db:
        assert len(db.scalars(select(OntologyEntity)).all()) == 6
        from app.models.relation import EntityRelation
        relations = db.scalars(select(EntityRelation)).all()
        assert all(rel.source_field is None and rel.target_field is None for rel in relations)
    rdf = b'''<rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#" xmlns:owl="http://www.w3.org/2002/07/owl#"><owl:Class rdf:about="https://example.test/#RdfDevice"/></rdf:RDF>'''
    rdf_preview = client.post('/entities/preview-file',data={'file_type':'rdf'},files={'file':('rdf.owl',rdf,'application/rdf+xml')})
    assert rdf_preview.status_code == 200,rdf_preview.text
    assert rdf_preview.json()['objects'][0]['name'] == 'RdfDevice'


def test_excel_preview_preserves_named_relation_columns_and_direction(environment):
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = '对象信息'
    ws.append(['对象英文名称','对象中文名称','属性英文名','属性中文名','属性类型','是否主键'])
    ws.append(['Device','设备','id','设备ID','string','是']); ws.append(['WorkOrder','工单','device_id','设备ID','string','是'])
    relations = wb.create_sheet('对象关系信息')
    relations.append(['对象','对象属性','被关联对象','被关联对象属性','关系标签1英文','基数设置'])
    relations.append(['WorkOrder','device_id','Device','id','for_device','多对一'])
    content = BytesIO(); wb.save(content); wb.close()
    client,_ = environment
    preview = client.post('/entities/preview-file',data={'file_type':'xlsx'},files={'file':('template.xlsx',content.getvalue(),'application/octet-stream')})
    assert preview.status_code == 200,preview.text
    relation = preview.json()['relations'][0]
    assert relation['cardinality'] == 'N:1'
    assert (relation['source_field'],relation['target_field']) == ('device_id','id')


def test_document_parser_preserves_late_content_and_word_tables(environment):
    from io import BytesIO
    from docx import Document
    from openpyxl import Workbook
    client,_ = environment
    content = ('设备需求\n'+'x'*9000+'\n最后必须生成工单').encode()
    response = client.post('/doc-builder/upload',files={'files':('large.md',content,'text/markdown')})
    assert response.status_code == 200,response.text
    key = f"doc-upload:{response.json()['session_id']}:large.md"
    result = client.get('/ai-builder/documents/content',params={'key':key}).json()
    assert result['truncated'] is True and result['total_chars'] > 9000
    doc = Document(); doc.add_paragraph('设备需求'); table = doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='风险阈值';table.cell(0,1).text='0.8'
    buffer=BytesIO();doc.save(buffer)
    assert '风险阈值\t0.8' in document_service.parse_file_content('device.docx',buffer.getvalue())
    wb=Workbook(); ws=wb.active;ws.cell(101,1,'最后一行业务约束'); buffer=BytesIO();wb.save(buffer);wb.close()
    assert '最后一行业务约束' in document_service.parse_file_content('device.xlsx',buffer.getvalue())


def test_document_chat_normalizes_validated_contract_and_invalidates_failed_result(environment, monkeypatch):
    import json
    from types import SimpleNamespace
    client,sessions = environment
    uploaded = client.post('/doc-builder/upload',files={'files':('device.md',b'Device id risk','text/markdown')}).json()['session_id']
    canonical = {'entities':[{'name':'Device','name_cn':'设备','tier':2,'description':'设备定义','attributes':[{'name':'id','display_name':'编号','type':'string','required':True,'description':'实例主键'},{'name':'risk','display_name':'风险','type':'number','required':False,'description':'风险分数'}]}],'relations':[]}
    class Completion:
        payload = canonical
        def create(self,**kwargs):
            text = json.dumps(self.payload,ensure_ascii=False)
            if kwargs.get('stream'):
                return iter([SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content=text[i:i+5]))]) for i in range(0,len(text),5)])
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text))])
    completion = Completion(); llm=SimpleNamespace(chat=SimpleNamespace(completions=completion))
    monkeypatch.setattr(document_service,'_get_llm_client',lambda db=None:llm)
    monkeypatch.setattr(document_service,'_get_model_name',lambda db=None:'isolated-fixture')
    response = client.post('/doc-builder/chat',json={'session_id':uploaded,'message':'','business_desc':'设备风险'})
    assert response.status_code == 200,response.text
    frames = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: {')]
    result = next(frame['data'] for frame in frames if frame['event']=='ontology')
    assert result['entities'][0]['displayName'] == '设备'
    assert result['entities'][0]['properties'][1]['name'] == 'risk'
    assert result['entities'][0]['tier'] == 2
    document_service._sessions.clear()
    assert client.get(f'/doc-builder/sessions/{uploaded}').json()['current_ontology'] == result
    completion.payload = {'entities':[],'relations':[]}
    failed = client.post('/doc-builder/chat',json={'session_id':uploaded,'message':'重新调整','business_desc':'设备风险'})
    assert '"event": "error"' in failed.text
    assert '"event": "ontology"' not in failed.text
    assert client.get(f'/doc-builder/sessions/{uploaded}').json()['current_ontology'] is None
    with sessions() as db:
        assert document_service.get_session(uploaded,db,'owner')['history'][-1]['role'] == 'assistant'


def test_document_apply_rejects_cross_ontology_and_rolls_back_new_definitions(environment):
    client, sessions = environment
    created = client.post('/builder/finalize', json=draft('ob')).json()['entity_ids']['Device']
    wrong = client.post('/doc-builder/mapping/apply', json={'session_id':'doc','ontology_id':'oa','items':[{'entity_id':created,'asset_id':'asset','field_mappings':[]}]})
    assert wrong.status_code == 422
    mapping = {'entities':[{'name':'NewDevice','asset_id':'asset','properties':[{'name':'id','field':'missing','type':'string'}]}], 'relations':[]}
    failed = client.post('/doc-builder/mapping/apply', json={'session_id':'doc','ontology_id':'oa','mapping_result':mapping,'items':[{'entity_name':'NewDevice','asset_id':'asset'}]})
    assert failed.status_code == 422
    with sessions() as db:
        assert db.scalar(select(OntologyEntity).where(OntologyEntity.name == 'NewDevice')) is None


def test_hydration_queries_real_source_and_reports_missing_fields_and_required_nulls(environment, tmp_path, monkeypatch):
    import sqlite3
    from app.config import settings
    from app.models.connection import Connection
    from app.services.builder.hydration_service import HydrationRequest, HydrationService
    monkeypatch.setattr(settings, 'LOCAL_DATA_ROOT', str(tmp_path))
    source = tmp_path / 'devices.db'
    with sqlite3.connect(source) as conn:
        conn.execute('CREATE TABLE devices (device_id TEXT PRIMARY KEY, risk_score REAL)')
        conn.executemany('INSERT INTO devices VALUES (?,?)', [('D1',0.8),('D2',0.2)])
    _, sessions = environment
    with sessions() as db:
        db.add(Connection(id='builder-real-source',name='builder-real-source',category='database',type='sqlite',database=str(source),params={'database':str(source)},credential_ref='',enabled=True))
        asset = db.get(Asset,'asset'); asset.connection_id = 'builder-real-source'; db.commit()
        request = {**draft('oa'),'session_id':'real-hydrate','asset_ids':[],'relations':[]}
        request.pop('ontology_id')
        request['objects'][0]['properties'][1]['required'] = True
        events = list(HydrationService(db).run(HydrationRequest(**request)))
        assert events[-1]['type'] == 'drill_complete'
        assert events[-1]['status'] == 'pass', events
        assert events[-1]['result']['entityCount'] == 2
        with sqlite3.connect(source) as conn:
            conn.execute("UPDATE devices SET risk_score=NULL WHERE device_id='D2'")
        failed = list(HydrationService(db).run(HydrationRequest(**request)))
        assert failed[-1]['status'] == 'error'
        assert any('必填字段存在空值' in event.get('msg','') for event in failed)
        request['objects'][0]['properties'][1]['source_column'] = 'missing'
        missing = list(HydrationService(db).run(HydrationRequest(**request)))
        assert missing[-1]['status'] == 'error'


def test_asset_and_document_extraction_use_explicit_asset_ids_with_deterministic_llm(environment, monkeypatch):
    import json
    from types import SimpleNamespace
    from app.services import ai_builder_v2, ontology_mapping_service, asset_catalog
    client, sessions = environment
    ontology = {'entities':[{'name':'Device','displayName':'设备','table':'asset','primaryKey':'id','properties':[{'name':'id','type':'string','source_field':'device_id'},{'name':'risk','type':'number','source_field':'risk_score'}]}],'relations':[]}
    class Completion:
        def create(self, **kwargs):
            prompt = kwargs['messages'][0]['content']
            if kwargs.get('stream'):
                text = json.dumps(ontology,ensure_ascii=False)
                return iter([SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content=text[i:i+7]))]) for i in range(0,len(text),7)])
            if 'candidates' in prompt:
                value = {'candidates':[{'entity':'Device','tables':['asset']}]}
            else:
                value = {'entities':[{'name':'Device','table':'asset','confidence':1,'properties':[{'name':'id','field':'device_id','confidence':1},{'name':'risk','field':'risk_score','confidence':1}]}],'relations':[]}
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(value)))])
    llm = SimpleNamespace(chat=SimpleNamespace(completions=Completion()))
    for module in (ai_builder_v2,ontology_mapping_service):
        monkeypatch.setattr(module,'_get_llm_client',lambda db=None:llm)
        monkeypatch.setattr(module,'_get_model_name',lambda db=None:'isolated-deterministic-fixture')
    with sessions() as db:
        db.add(Asset(id='asset-other',name='other',kind='table',status='active',locator={'table':'devices'},schema_snapshot=[{'name':'other_id','type':'TEXT'}])); db.commit()
        with pytest.raises(ValueError,match='多个资产'):
            asset_catalog.get_table_schema('devices',db)
        events = list(ai_builder_v2.extract_ontology_stream(['asset'],[],'设备风险',db))
        result = next(json.loads(line[6:])['data'] for line in events if '"event": "result"' in line)
        assert result['entities'][0]['backing_asset_ids'] == ['asset']
        assert result['entities'][0]['properties'][0]['source_column'] == 'device_id'
        assert result['entities'][0]['properties'][1]['source_asset_id'] == 'asset'
        mapped_events = list(ontology_mapping_service.map_ontology_stream(ontology,db))
        mapped = next(json.loads(line[6:])['data'] for line in mapped_events if '"event": "result"' in line)
        assert mapped['entities'][0]['table'] == 'asset'
        assert mapped['entities'][0]['properties'][1]['field'] == 'risk_score'
    # Through the actual apply API; only LLM generation above is substituted.
    applied = client.post('/doc-builder/mapping/apply',json={'session_id':'isolated','ontology_id':'oa','mapping_result':mapped,'items':[{'entity_name':'Device','asset_id':'asset'}]})
    assert applied.status_code == 200, applied.text
    assert applied.json()['created'] == 1


@pytest.mark.parametrize('entry', ['manual', 'template', 'document', 'asset'])
def test_four_builder_entries_publish_and_query_real_instances(environment, tmp_path, monkeypatch, entry):
    """Each supported entry goes through real HTTP draft/binding/snapshot/SQLite readback.

    Only LLM completions are deterministic substitutes. Source rows are allowed
    fixture inputs; no entities, bindings or version success states are seeded.
    """
    import json
    import sqlite3
    from io import BytesIO
    from types import SimpleNamespace
    from openpyxl import Workbook
    from app.config import settings
    from app.models.connection import Connection
    from app.models.version import OntologyVersion, OntologyVersionEntity
    from app.services import ai_builder_v2 as asset_service, ontology_mapping_service as mapping_service
    client, sessions = environment
    monkeypatch.setattr(settings, 'LOCAL_DATA_ROOT', str(tmp_path))
    source = tmp_path / f'{entry}.db'
    with sqlite3.connect(source) as conn:
        conn.execute('CREATE TABLE devices (device_id TEXT PRIMARY KEY, risk_score REAL)')
        conn.executemany('INSERT INTO devices VALUES (?, ?)', [('DEV-001',0.92),('DEV-002',0.15)])
    # Connection + source Asset are prerequisites only; the tested business
    # operations below must create all ontology structures and published state.
    with sessions() as db:
        db.add(Connection(id=f'{entry}-connection',name=entry,category='database',type='sqlite',database=str(source),params={'database':str(source)},credential_ref='',enabled=True))
        db.get(Asset,'asset').connection_id = f'{entry}-connection'; db.commit()
    scenario = client.post('/scenarios',json={'code':f'entry-{entry}','name':f'{entry} 验收'})
    assert scenario.status_code == 201, scenario.text
    ontology_id = scenario.json()['id']
    canonical = {'entities':[{'name':'Device','name_cn':'设备','tier':2,'description':'设备风险','attributes':[{'name':'id','display_name':'编号','type':'string','required':True,'description':'主键'},{'name':'risk','display_name':'风险','type':'number','required':True,'description':'风险值'}]}],'relations':[]}
    extracted = {'entities':[{'name':'Device','displayName':'设备','description':'设备风险','tier':2,'table':'asset','primaryKey':'id','properties':[{'name':'id','type':'string','required':True,'source_field':'device_id'},{'name':'risk','type':'number','required':True,'source_field':'risk_score'}]}],'relations':[]}
    precise = {'entities':[{'name':'Device','table':'asset','confidence':1,'properties':[{'name':'id','field':'device_id','confidence':1},{'name':'risk','field':'risk_score','confidence':1}]}],'relations':[]}
    class Completion:
        def __init__(self, document=False): self.document = document
        def create(self, **kwargs):
            prompt = kwargs['messages'][0]['content']
            value = canonical if self.document and kwargs.get('stream') else extracted if kwargs.get('stream') else {'candidates':[{'entity':'Device','tables':['asset']}]} if 'candidates' in prompt else precise
            text = json.dumps(value, ensure_ascii=False)
            if kwargs.get('stream'):
                return iter([SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content=text[i:i+3]))]) for i in range(0,len(text),3)])
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text))])
    for service in (asset_service,mapping_service,document_service):
        client_fixture = SimpleNamespace(chat=SimpleNamespace(completions=Completion(service is document_service)))
        monkeypatch.setattr(service,'_get_llm_client',lambda db=None, llm=client_fixture:llm)
        monkeypatch.setattr(service,'_get_model_name',lambda db=None:'controlled-completion')
    def frames(response):
        assert response.status_code == 200, response.text
        assert 'data: [DONE]' in response.text
        result = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: {')]
        assert not any(frame.get('event') == 'error' or frame.get('type') == 'error' for frame in result), result
        return result
    payload = draft(ontology_id)
    payload['objects'][0]['properties'][1]['required'] = True
    if entry == 'template':
        wb=Workbook(); ws=wb.active; ws.title='对象信息'
        ws.append(['对象英文名称','对象中文名称','属性英文名','属性中文名','属性类型','是否主键'])
        ws.append(['Device','设备','id','编号','string','是']); ws.append(['Device','设备','risk','风险','number','否'])
        buffer=BytesIO();wb.save(buffer);wb.close()
        preview=client.post('/entities/preview-file',data={'file_type':'xlsx'},files={'file':('device.xlsx',buffer.getvalue(),'application/octet-stream')})
        assert preview.status_code == 200,preview.text
        obj=preview.json()['objects'][0]
        assert obj['primary_key'] == 'id'
        payload['objects'][0].update({'name':obj['name'],'displayName':obj['display_name'],'primaryKey':obj['primary_key']})
        payload['objects'][0]['properties']=[{'name':p['name'],'type':p['type'],'required':p['name']=='id','source_asset_id':'asset','source_column':'device_id' if p['name']=='id' else 'risk_score'} for p in obj['properties']]
    elif entry == 'asset':
        result=next(frame['data'] for frame in frames(client.post('/ai-builder/extract-ontology',json={'table_names':['asset'],'business_desc':'设备风险','document_keys':[]})) if frame.get('event') == 'result')
        payload['objects'] = result['entities']
        assert payload['objects'][0]['backing_asset_ids'] == ['asset']
    elif entry == 'document':
        uploaded=client.post('/doc-builder/upload',files={'files':('device.md',b'Device id risk','text/markdown')})
        assert uploaded.status_code == 200,uploaded.text
        session_id=uploaded.json()['session_id']
        ontology=next(frame['data'] for frame in frames(client.post('/doc-builder/chat',json={'session_id':session_id,'business_desc':'设备风险','message':''})) if frame.get('event')=='ontology')
        ontology['entities'][0]['primaryKey'] = 'id' # User review chooses a defined property.
        mapped=next(frame['data'] for frame in frames(client.post('/doc-builder/mapping',json={'session_id':session_id,'ontology':ontology})) if frame.get('event')=='result')
        mapped['entities'][0]['asset_id'] = mapped['entities'][0]['table']
        mapped['entities'][0]['table'] = 'devices' # UI selector ID to physical table conversion.
        preview=client.post('/doc-builder/mapping/preview',json={'session_id':session_id,'ontology_id':ontology_id,'mapping_result':mapped})
        assert preview.status_code == 200,preview.text
        item=preview.json()['items'][0]
        applied=client.post('/doc-builder/mapping/apply',json={'session_id':session_id,'ontology_id':ontology_id,'mapping_result':mapped,'items':[{'entity_name':item['entity_name'],'entity_id':item['entity_id'],'asset_id':item['asset_id']}]})
        assert applied.status_code == 200,applied.text
        assert applied.json()['created'] == 1
    hydrated=frames(client.post('/builder/hydrate',json={**payload,'session_id':f'{entry}-session','asset_ids':['asset'],'relations':[]}))[-1]
    assert hydrated['type'] == 'drill_complete' and hydrated['status'] == 'pass',hydrated
    assert hydrated['result']['entityCount'] == 2
    finalized=client.post('/builder/finalize',json=payload)
    assert finalized.status_code == 200,finalized.text
    entity_id=finalized.json()['entity_ids']['Device']
    # Fresh ORM session prevents passing due to objects left in the caller cache.
    with sessions() as db:
        binding=db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_id))
        assert binding.asset_id == 'asset' and binding.id_column == 'device_id'
        attributes=db.scalars(select(EntityAttribute).where(EntityAttribute.entity_id==entity_id)).all()
        assert {field['attribute_id'] for field in binding.field_mappings} == {attr.id for attr in attributes}
        binding_id = binding.id
    published=client.post('/ontology-publish/quick-publish',json={'ontology_id':ontology_id,'name':f'{entry} v1'})
    assert published.status_code == 200,published.text
    version_id=published.json()['version_id']
    queried=client.post('/ontology-api/query',json={'entity_name':'Device','ontology_id':ontology_id,'version_id':version_id,'filters':{'id':'DEV-001'},'depth':0})
    assert queried.status_code == 200,queried.text
    assert queried.json()['instances'] == [{'id':'DEV-001','risk':0.92}]
    with sessions() as db:
        version=db.get(OntologyVersion,version_id)
        assert version.status == 'published' and version.is_active is True
        snapshot=db.scalar(select(OntologyVersionEntity).where(OntologyVersionEntity.version_id==version_id))
        assert snapshot.binding_snapshot and snapshot.source_entity_id == entity_id
    # Repeat the consumption request after another independent session.
    assert client.get('/ontology-api/objects/Device/DEV-002',params={'ontology_id':ontology_id,'version_id':version_id}).json() == {'id':'DEV-002','risk':0.15}
    print(json.dumps({'flow':'B','entry':entry,'ontology_id':ontology_id,'entity_id':entity_id,'binding_id':binding_id,'version_id':version_id,'source':'SQLite temp file','readback':queried.json(),'llm_substitution':entry in ['document','asset']},ensure_ascii=False))


def test_studio_json_roundtrip_preserves_fields_and_real_template_actions(environment):
    import json
    from app.models.action import EntityAction
    from app.models.relation import EntityRelation
    client,sessions=environment
    exported={'objects':[{'name':'Device','display_name':'设备','tier':2,'primary_key':'id','properties':[{'name':'id','type':'string','required':True,'source_table':'devices','source_field':'device_id'},{'name':'risk','type':'number','source_table':'devices','source_field':'risk_score'}]},{'name':'OtherDevice','display_name':'另设备','tier':2,'primary_key':'id','properties':[{'name':'id','type':'string','required':True,'source_table':'devices','source_field':'device_id'}]}], 'relations':[{'name':'related','source':'Device','target':'OtherDevice','cardinality':'N:1','source_field':'id','target_field':'id'}], 'actions':[{'name':'DeviceRisk','display_name':'风险返回','target_object':'Device','action_type':'custom_script','type_config':{'script':'result = {"risk": params["risk"] * 100}'},'parameters_json':[{'name':'risk','type':'number','required':True}],'description':'确定性风险返回'},{'name':'DeviceRepairSuggestion','target_object':'Device','trigger':'automatic','description':'高风险时建议检修','effects':[{'target':'Device','attribute':'status','value':'repair'}]}]}
    content=json.dumps(exported,ensure_ascii=False).encode()
    preview=client.post('/entities/preview-file',data={'file_type':'json'},files={'file':('export.json',content,'application/json')})
    assert preview.status_code == 200,preview.text
    assert preview.json()['summary']['object_count'] == 2
    assert preview.json()['relations'][0]['source_field'] == 'id'
    assert preview.json()['actions'][0]['type_config']['script'].startswith('result =')
    saved=client.post('/entities/from-file',data={'file_type':'json','ontology_id':'oa'},files={'file':('export.json',content,'application/json')})
    assert saved.status_code == 201,saved.text
    assert saved.json()['actions_created'] == 1
    assert any('待配置建议' in warning for warning in saved.json()['warnings'])
    with sessions() as db:
        assert db.scalar(select(EntityRelation)).source_field == 'id'
        assert len(db.scalars(select(ObjectBinding)).all()) == 2
        entity=db.scalar(select(OntologyEntity).where(OntologyEntity.ontology_id=='oa',OntologyEntity.name=='Device'))
        assert entity.config_json['pending_action_specs'][0]['name'] == 'DeviceRepairSuggestion'
        action=db.scalar(select(EntityAction))
        from app.services.action_executors import run_executor_sync
        actual=run_executor_sync(action.action_type,action.type_config,{'risk':0.92},db=db)
        assert actual.success is True and actual.output == {'risk':92.0}
    # Template finalize consumes the same structured preview rather than losing
    # executable configuration in the old object.actions name-only array.
    request=draft('ob'); request['actions']=preview.json()['actions']
    result=client.post('/builder/finalize',json=request)
    assert result.status_code == 200,result.text
    assert result.json()['created_actions'] == 1
    assert client.post('/builder/finalize',json=request).json()['created_actions'] == 0
    with sessions() as db:
        assert len(db.scalars(select(EntityAction)).all()) == 2


def test_finalize_replaces_removed_fields_and_source_bindings_without_touching_other_ontology(environment):
    from app.models.quality_rule import QualityRule
    client,sessions=environment
    entity_a=client.post('/builder/finalize',json=draft('oa')).json()['entity_ids']['Device']
    entity_b=client.post('/builder/finalize',json=draft('ob')).json()['entity_ids']['Device']
    with sessions() as db:
        previous=db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_a)).id
        db.add(Asset(id='replacement',name='replacement',kind='table',status='active',locator={'table':'replacement'},schema_snapshot=[{'name':'new_id','type':'string'}],primary_key=['new_id']));db.commit()
    changed={'ontology_id':'oa','objects':[{'name':'Device','displayName':'设备','primaryKey':'id','backing_asset_ids':['replacement'],'properties':[{'name':'id','type':'string','required':True,'source_asset_id':'replacement','source_column':'new_id'}]}]}
    result=client.post('/builder/finalize',json=changed)
    assert result.status_code == 200,result.text
    assert result.json()['entity_ids']['Device'] == entity_a
    with sessions() as db:
        assert db.get(ObjectBinding,previous) is None
        bindings=db.scalars(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_a)).all()
        assert len(bindings)==1 and bindings[0].asset_id=='replacement' and bindings[0].id_column=='new_id'
        assert [attribute.name for attribute in db.scalars(select(EntityAttribute).where(EntityAttribute.entity_id==entity_a))] == ['id']
        assert not db.scalars(select(QualityRule).where(QualityRule.description==f'auto:binding:{previous}')).all()
        assert db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_b)).asset_id=='asset'
        assert len(db.scalars(select(EntityAttribute).where(EntityAttribute.entity_id==entity_b)).all())==2


def test_asset_extraction_and_mapping_failures_end_truthfully(environment, monkeypatch):
    import json
    from types import SimpleNamespace
    from app.services import ai_builder_v2 as asset_service, ontology_mapping_service as mapping_service
    client,_=environment
    class Completion:
        def create(self,**kwargs):
            if kwargs.get('stream'):
                return iter([SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content='{"entities":[],"relations":[]}'))])])
            raise ValueError('LLM downstream failed')
    llm=SimpleNamespace(chat=SimpleNamespace(completions=Completion()))
    for service in (asset_service,mapping_service):
        monkeypatch.setattr(service,'_get_llm_client',lambda db=None:llm)
        monkeypatch.setattr(service,'_get_model_name',lambda db=None:'isolated-failure')
    for endpoint,payload in [('/ai-builder/extract-ontology',{'table_names':['asset'],'document_keys':[],'business_desc':'risk'}),('/doc-builder/mapping',{'session_id':'failure','ontology':{'entities':[{'name':'Device','properties':[{'name':'id'}]}],'relations':[]}})]:
        result=client.post(endpoint,json=payload)
        assert result.status_code==200
        assert result.text.endswith('data: [DONE]\n\n')
        events=[json.loads(line[6:]) for line in result.text.splitlines() if line.startswith('data: {')]
        assert any(frame.get('event')=='error' for frame in events)
        assert not any(frame.get('event')=='result' for frame in events)


def test_executable_template_actions_require_rule_write_on_both_import_routes(environment):
    import json
    from app.models.system_config import SystemConfig
    client,sessions=environment
    with sessions() as db:
        db.add(SystemConfig(key='auth.role_permissions.editor',value=json.dumps(['entity:read','entity:write']),group='auth',description='editor'));db.commit()
    client.app.dependency_overrides[require_user]=lambda:User(id='entity-writer',username='writer',name='Writer',password_hash='x',role='editor',is_active=True)
    action={'name':'risk','target_object':'Device','action_type':'custom_script','type_config':{'script':'result = {"risk": 1}'}}
    request=draft('oa'); request['actions']=[action]
    assert client.post('/builder/finalize',json=request).status_code == 403
    imported={'objects':[{'name':'Device','properties':[{'name':'id','type':'string'}]}],'actions':[action]}
    assert client.post('/entities/from-file',data={'file_type':'json','ontology_id':'oa'},files={'file':('actions.json',json.dumps(imported).encode(),'application/json')}).status_code == 403
    with sessions() as db:
        assert db.scalar(select(OntologyEntity)) is None


def test_hydration_verifies_cross_connection_real_relation_keys(environment, tmp_path, monkeypatch):
    import json
    import sqlite3
    from app.config import settings
    from app.models.connection import Connection
    client,sessions=environment
    monkeypatch.setattr(settings,'LOCAL_DATA_ROOT',str(tmp_path))
    source=tmp_path/'source.db'; target=tmp_path/'target.db'
    with sqlite3.connect(source) as conn:
        conn.execute('CREATE TABLE devices (device_id TEXT PRIMARY KEY, risk_score REAL)');conn.execute('INSERT INTO devices VALUES (?,?)',('DEV-001',0.92))
    with sqlite3.connect(target) as conn:
        conn.execute('CREATE TABLE orders (order_id TEXT PRIMARY KEY, device_id TEXT NOT NULL)');conn.executemany('INSERT INTO orders VALUES (?,?)',[('WO-1','DEV-001'),('WO-2','DEV-001'),('WO-3','DEV-002')])
    with sessions() as db:
        for cid,path in [('source',source),('target',target)]:
            db.add(Connection(id=cid,name=cid,category='database',type='sqlite',params={'database':str(path)},database=str(path),credential_ref='',enabled=True))
        db.get(Asset,'asset').connection_id='source'
        db.add(Asset(id='orders',name='orders',kind='table',status='active',connection_id='target',locator={'table':'orders'},schema_snapshot=[{'name':'order_id','type':'TEXT'},{'name':'device_id','type':'TEXT'}],primary_key=['order_id']));db.commit()
    objects=draft('oa')['objects']+[{'name':'WorkOrder','primaryKey':'id','backing_asset_ids':['orders'],'properties':[{'name':'id','type':'string','required':True,'source_asset_id':'orders','source_column':'order_id'},{'name':'device','type':'string','required':True,'source_asset_id':'orders','source_column':'device_id'}]}]
    request={'session_id':'cross-connection','asset_ids':[],'objects':objects,'relations':[{'name':'orders','source':'Device','target':'WorkOrder','sourceField':'id','targetField':'device','cardinality':'1:N'}]}
    def terminal():
        response=client.post('/builder/hydrate',json=request)
        frames=[json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: {')]
        return frames,frames[-1]
    events,last=terminal()
    assert last['status']=='pass',events
    assert last['result']['relationCount']==2 and last['result']['entityCount']==4
    assert any('跨连接真实关联验证通过' in event.get('msg','') for event in events)
    with sessions() as db:
        db.get(Asset,'orders').sensitivity_tags={'device_id':'pii'};db.commit()
    events,last=terminal()
    assert last['status']=='error'
    assert any('敏感关联键' in event.get('msg','') for event in events)
    request['relations'][0]['sourceField']=None
    events,last=terminal()
    assert last['status']=='error'
    assert any('未确认源/目标关联字段' in event.get('msg','') for event in events)


def test_action_receipt_reloads_same_account_without_redispatch(environment):
    from app.models.action import EntityAction
    from app.models.action_execution import ActionExecution
    client,sessions=environment
    created=client.post('/actions',json={'name':'Risk','category':'system','ontology_id':'oa','action_type':'custom_script',
        'type_config':{'script':'result = {"risk": params["temperature"] / 100}'},
        'parameters_json':[{'name':'temperature','type':'number','required':True}]})
    assert created.status_code==201,created.text
    action_id=created.json()['id']
    request={'params':{'temperature':92},'idempotency_key':'ui-owner-reload','dry_run':False}
    first=client.post(f'/actions/{action_id}/execute',json=request)
    assert first.status_code==200 and first.json()['success'],first.text
    assert first.json()['output']['risk']==0.92
    repeat=client.post(f'/actions/{action_id}/execute',json=request)
    assert repeat.json()['execution_id']==first.json()['execution_id']
    receipt=client.get(f'/actions/{action_id}/executions/by-key',params={'key':'ui-owner-reload'})
    assert receipt.status_code==200 and receipt.json()['state']=='succeeded',receipt.text
    assert receipt.json()['params']=={'temperature':92}
    assert receipt.json()['execution_id']==first.json()['execution_id']
    with sessions() as db:
        assert len(db.scalars(select(ActionExecution)).all())==1
        assert db.get(EntityAction,action_id).impact_count==1
    client.app.dependency_overrides[require_user]=lambda:User(id='other',username='other',password_hash='x',role='admin',is_active=True)
    assert client.get(f'/actions/{action_id}/executions/by-key',params={'key':'ui-owner-reload'}).status_code==404


def test_ai_code_modes_execute_generated_expression_and_isolate_private_history(environment,monkeypatch):
    import json
    from types import SimpleNamespace
    from app.services import ai_code_service
    from app.models.ai_code_conversation import AiCodeConversation
    client,sessions=environment
    captured=[]
    expression="params['temperature'] / 100"
    def completion(**kwargs):
        captured.append(kwargs['messages'])
        mode=kwargs['messages'][0]['content']
        text='result = {"risk": params["temperature"] / 100}' if 'custom_script' in mode else ('SELECT device_id FROM devices WHERE risk_score > :risk' if '当前执行模式：sql' in mode else expression)
        return iter([SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content=text))])])
    monkeypatch.setattr(ai_code_service,'get_llm_client',lambda *a,**k:SimpleNamespace(_model_name='isolated-controlled',chat=SimpleNamespace(completions=SimpleNamespace(create=completion))))
    stream=client.post('/ai-code/generate',json={'target_type':'function','target_id':'tmp_generated','message':'风险=温度/100','extra_entity_ids':[],'execution_mode':'expression'})
    assert '"event": "done"' in stream.text and expression in stream.text,stream.text
    assert '不允许赋值' in captured[-1][0]['content']
    assert client.post('/ai-code/validate',json={'code':expression,'execution_mode':'expression'}).json()['safe'] is True
    assert client.post('/ai-code/validate',json={'code':'result = 1','execution_mode':'expression'}).json()['safe'] is False
    function=client.post('/functions',json={'name':'GeneratedRisk','callable_name':'generated_risk','ontology_id':'oa','return_type':'number','logic_type':'expression','logic_body':expression,'input_schema':[{'name':'temperature','type':'number','required':True}]})
    assert function.status_code==201,function.text
    ran=client.post('/functions/'+function.json()['id']+'/test',json={'params':{'temperature':92}})
    assert ran.status_code==200 and ran.json()['result']==0.92,ran.text
    action=client.post('/actions',json={'name':'GeneratedAction','category':'system','ontology_id':'oa','action_type':'custom_script','type_config':{'script':'result = {"risk": 0}'}}).json()
    generated=client.post('/ai-code/generate',json={'target_type':'action','target_id':action['id'],'message':'风险按输入温度除100','execution_mode':'custom_script'})
    assert '"event": "done"' in generated.text and 'GeneratedAction' in captured[-1][0]['content'],generated.text
    client.post('/builder/finalize',json=draft('oa')).raise_for_status()
    entity_id=client.get('/entities',params={'ontology_id':'oa'}).json()[0]['id']
    sql=client.post('/ai-code/generate',json={'target_type':'function','target_id':'tmp_sql','message':'查风险设备','extra_entity_ids':[entity_id],'execution_mode':'sql'})
    assert 'SELECT device_id' in sql.text and 'devices' in captured[-1][0]['content']
    assert client.post('/ai-code/validate',json={'code':'SELECT device_id FROM devices WHERE risk_score > :risk','execution_mode':'sql'}).json()['safe'] is True
    assert client.post('/ai-code/validate',json={'code':'DELETE FROM devices','execution_mode':'sql'}).json()['safe'] is False
    own=client.get('/ai-code/conversations/function/tmp_generated').json()
    assert own['messages'][-1]['content']==expression
    with sessions() as db:
        assert db.scalar(select(AiCodeConversation).where(AiCodeConversation.target_id=='tmp_generated')).owner_id=='owner'
    client.app.dependency_overrides[require_user]=lambda:User(id='other',username='other',password_hash='x',role='admin',is_active=True)
    assert client.get('/ai-code/conversations/function/tmp_generated').json() is None
    # The other account starts a distinct private history even for the same target.
    other=client.post('/ai-code/generate',json={'target_type':'function','target_id':'tmp_generated','message':'另一个会话','execution_mode':'expression'})
    assert '"event": "done"' in other.text
    assert client.get('/ai-code/conversations/function/tmp_generated').json()['id']!=own['id']


def test_ai_private_history_upgrade_preserves_unowned_legacy_rows():
    from sqlalchemy import text, inspect
    from app.migrations import run_startup_migrations
    engine=create_engine('sqlite://')
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE ai_code_conversations (id VARCHAR(36) PRIMARY KEY, target_type VARCHAR(20), target_id VARCHAR(36), messages JSON, context_entity_ids JSON, created_at DATETIME, updated_at DATETIME)"))
        connection.execute(text("INSERT INTO ai_code_conversations(id,target_type,target_id,messages,context_entity_ids) VALUES ('legacy1','function','same','[]','[]'), ('legacy2','function','same','[]','[]')"))
    run_startup_migrations(engine)
    run_startup_migrations(engine)
    assert 'owner_id' in {column['name'] for column in inspect(engine).get_columns('ai_code_conversations')}
    with engine.connect() as connection:
        assert connection.execute(text('SELECT id, owner_id FROM ai_code_conversations ORDER BY id')).all()==[('legacy1',None),('legacy2',None)]
    from sqlalchemy.orm import Session
    from app.services.ai_code_service import AiCodeService
    with Session(engine) as db:
        assert AiCodeService(db,'new-owner').get_conversation('function','same') is None
        private=AiCodeService(db,'new-owner')._get_or_create_conversation('function','same',[])
        db.commit()
        assert AiCodeService(db,'new-owner').get_conversation('function','same').id==private.id
        assert AiCodeService(db,'other-owner').get_conversation('function','same') is None
    engine.dispose()


def test_auto_mapping_uses_its_own_request_contract_and_finalize_requires_primary(environment):
    client, sessions = environment
    response=client.post('/builder/auto-map',json={'objects':[{'id':'o','name':'devices','properties':[{'name':'device_id','type':'string'},{'name':'risk_score','type':'number'}]}],'asset_ids':['asset']})
    assert response.status_code == 200,response.text
    assert response.json()['mappings']['o']['asset_id'] == 'asset'
    payload=draft('oa');payload['objects'][0]['primaryKey']='missing'
    assert client.post('/builder/finalize',json=payload).status_code == 422
    with sessions() as db: assert db.scalar(select(OntologyEntity)) is None


def test_builder_enrichment_requires_explicit_rules_and_queries_published_left_join(environment,tmp_path,monkeypatch):
    import json
    import sqlite3
    from app.config import settings
    from app.models.connection import Connection
    from app.services.data_plane.execute_service import ExecuteService
    client,sessions=environment
    monkeypatch.setattr(settings,'LOCAL_DATA_ROOT',str(tmp_path))
    primary,extra=tmp_path/'devices.db',tmp_path/'owners.db'
    with sqlite3.connect(primary) as raw:
        raw.execute('CREATE TABLE devices(device_id TEXT PRIMARY KEY,risk_score REAL)')
        raw.executemany('INSERT INTO devices VALUES (?,?)',[('DEV-001',.92),('DEV-002',.15)])
    with sqlite3.connect(extra) as raw:
        raw.execute('CREATE TABLE owners(record_id TEXT PRIMARY KEY,serial TEXT,owner TEXT)')
        raw.execute("INSERT INTO owners VALUES ('REC1','DEV-001','Team A')")
    with sessions() as db:
        for key,file in [('primary',primary),('extra',extra)]:
            db.add(Connection(id=key,name=key,category='database',type='sqlite',database=str(file),params={'database':str(file)},credential_ref='',enabled=True))
        db.get(Asset,'asset').connection_id='primary'
        db.add(Asset(id='owners',name='owners',kind='table',connection_id='extra',status='active',locator={'table':'owners'},primary_key=['record_id'],schema_snapshot=[{'name':key,'type':'string'} for key in ['record_id','serial','owner']]))
        db.commit()
    payload=draft('oa');obj=payload['objects'][0]
    obj['backing_asset_ids'].append('owners')
    obj['properties'].append({'name':'owner','type':'string','source_asset_id':'owners','source_column':'owner'})
    def hydrate():
        response=client.post('/builder/hydrate',json={**payload,'session_id':'enriched','asset_ids':['asset','owners'],'relations':[]})
        assert response.status_code==200,response.text
        return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: {')][-1]
    assert hydrate()['status']=='error'
    initial=client.post('/builder/finalize',json=payload)
    assert initial.status_code==200,initial.text
    entity_id=initial.json()['entity_ids']['Device']
    with sessions() as db:
        binding=db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_id,ObjectBinding.role=='enrichment'))
        assert binding.status=='needs_review' and binding.primary_source_column is None
    blocked=client.post('/ontology-publish/quick-publish',json={'ontology_id':'oa','name':'Invalid join'})
    assert blocked.status_code==400,blocked.text
    assert '一致性检查未通过' in blocked.json()['detail'],blocked.text
    from app.models.version import OntologyVersion
    with sessions() as db: assert db.scalar(select(OntologyVersion).where(OntologyVersion.ontology_id=='oa')) is None
    obj['enrichment_bindings']=[{'asset_id':'owners','primary_source_column':'device_id','enrichment_source_column':'serial'}]
    assert hydrate()['status']=='pass'
    saved=client.post('/builder/finalize',json=payload);assert saved.status_code==200,saved.text
    with sessions() as db:
        binding=db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_id,ObjectBinding.role=='enrichment'))
        assert binding.status=='active' and binding.enrichment_source_column=='serial'
    published=client.post('/ontology-publish/quick-publish',json={'ontology_id':'oa','name':'Explicit join'})
    assert published.status_code==200,published.text
    response=client.post('/ontology-api/query',json={'entity_name':'Device','ontology_id':'oa','version_id':published.json()['version_id'],'depth':0})
    assert response.status_code==200,response.text
    assert response.json()['instances']==[{'id':'DEV-001','risk':.92,'owner':'Team A'},{'id':'DEV-002','risk':.15,'owner':None}]
    with sqlite3.connect(extra) as raw: raw.execute("INSERT INTO owners VALUES ('REC2','DEV-001','Duplicate')")
    ExecuteService.invalidate_cache_for_asset('owners')
    failed=hydrate();assert failed['status']=='error',failed


def test_builder_transform_is_preserved_and_hydration_uses_real_converted_values(environment,tmp_path,monkeypatch):
    import json
    import sqlite3
    from app.config import settings
    from app.models.connection import Connection
    client,sessions=environment
    monkeypatch.setattr(settings,'LOCAL_DATA_ROOT',str(tmp_path))
    source=tmp_path/'transformed.db'
    with sqlite3.connect(source) as raw:
        raw.execute('CREATE TABLE devices(device_id TEXT PRIMARY KEY,risk_score REAL)')
        raw.executemany('INSERT INTO devices VALUES (?,?)',[('DEV-001',.92),('DEV-002',None)])
    with sessions() as db:
        db.add(Connection(id='transform',name='transform',category='database',type='sqlite',database=str(source),params={'database':str(source)},credential_ref='',enabled=True))
        db.get(Asset,'asset').connection_id='transform';db.commit()
    payload=draft('oa');properties=payload['objects'][0]['properties']
    properties[1].update(required=True,transform='round((value or 0) * 100)')
    def hydrate():
        response=client.post('/builder/hydrate',json={**payload,'session_id':'transform','asset_ids':['asset'],'relations':[]})
        assert response.status_code==200,response.text
        return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: {')][-1]
    assert hydrate()['status']=='pass'
    saved=client.post('/builder/finalize',json=payload);assert saved.status_code==200,saved.text
    entity_id=saved.json()['entity_ids']['Device']
    with sessions() as db:
        binding=db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_id))
        assert binding.field_mappings[1]['transform']=='round((value or 0) * 100)'
    published=client.post('/ontology-publish/quick-publish',json={'ontology_id':'oa','name':'Transform'})
    assert published.status_code==200,published.text
    query=client.post('/ontology-api/query',json={'ontology_id':'oa','version_id':published.json()['version_id'],'entity_name':'Device','depth':0})
    assert query.status_code==200,query.text
    assert query.json()['instances']==[{'id':'DEV-001','risk':92},{'id':'DEV-002','risk':0}]
    properties[1]['transform']='value / 0'
    assert hydrate()['status']=='error'
    properties[1]['transform']='__import__("os")'
    assert client.post('/builder/finalize',json=payload).status_code==422
    properties[1]['transform']=None;properties[0]['transform']='str(value)'
    assert hydrate()['status']=='error'
    assert client.post('/builder/finalize',json=payload).status_code==422
    with sessions() as db:
        binding=db.scalar(select(ObjectBinding).where(ObjectBinding.object_type_id==entity_id))
        assert binding.field_mappings[1]['transform']=='round((value or 0) * 100)'
