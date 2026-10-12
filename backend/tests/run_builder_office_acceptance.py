"""Actual legacy Office conversion acceptance; only disposable local files are used.

No DB/app/LLM is seeded. Requires soffice on PATH or OFFICE_CONVERTER.
Run: SECRET_KEY=<isolated-key> python3 tests/run_builder_office_acceptance.py
"""
import json
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from app.config import settings
from app.services.doc_builder import parse_file_content
from app.services.excel_import import preview_excel_ontology
from app.services.office_conversion import OLE_MAGIC


def main():
    executable = settings.OFFICE_CONVERTER or shutil.which('soffice')
    if not executable:
        raise RuntimeError('LibreOffice missing: this real conversion acceptance cannot run')
    records=[]
    with tempfile.TemporaryDirectory(prefix='ontology-office-acceptance-') as directory:
        root=Path(directory)
        doc=Document();doc.add_paragraph('Device alarm risk'); doc.save(root/'source.docx')
        wb=Workbook();ws=wb.active;ws.title='对象信息'
        ws.append(['对象英文名称','对象中文名称','属性英文名','属性中文名','属性类型','是否主键'])
        ws.append(['Device','设备','id','设备ID','string','是']);wb.save(root/'source.xlsx');wb.close()
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='Device alarm risk';ppt.save(root/'source.pptx')
        for original,legacy,expected in [('docx','doc','Device alarm risk'),('xlsx','xls','Device'),('pptx','ppt','Device alarm risk')]:
            process=subprocess.run([executable,f'-env:UserInstallation={(root/(legacy+"-profile")).as_uri()}','--headless','--convert-to',legacy,'--outdir',str(root),str(root/f'source.{original}')],capture_output=True,timeout=60)
            target=root/f'source.{legacy}'
            assert process.returncode==0 and target.is_file(),process.stderr.decode(errors='replace')
            content=target.read_bytes();assert content.startswith(OLE_MAGIC)
            text=parse_file_content(f'source.{legacy}',content);assert expected in text
            entry={'format':legacy,'source_bytes':len(content),'expected_text':expected,'readback':text,'result':'passed'}
            if legacy=='xls':
                preview=preview_excel_ontology(content)
                assert preview['objects'][0]['name']=='Device'
                assert preview['objects'][0]['primary_key']=='id'
                entry['template_preview']=preview
            records.append(entry)
    print(json.dumps({'level':'real internal Office subprocess integration','source':'generated Office documents, temporary directory; no DB or network','converter':Path(executable).name,'records':records},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
