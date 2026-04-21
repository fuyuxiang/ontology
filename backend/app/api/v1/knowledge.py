import os
import uuid as _uuid
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from sqlalchemy.orm import Session
from sqlalchemy import or_

from app.database import get_db
from app.models.knowledge import KnowledgeBase, KnowledgeFile

router = APIRouter(prefix="/knowledge", tags=["knowledge"])

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "..", "uploads", "knowledge")
os.makedirs(UPLOAD_DIR, exist_ok=True)

EXT_TYPE_MAP = {
    "pdf": "pdf", "doc": "word", "docx": "word",
    "xls": "excel", "xlsx": "excel",
    "png": "image", "jpg": "image", "jpeg": "image", "gif": "image", "webp": "image",
    "mp4": "video", "avi": "video", "mov": "video", "mkv": "video",
}


def _parse_file(path: str, file_type: str) -> str:
    try:
        if file_type == "pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif file_type == "word":
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif file_type == "excel":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)
    except Exception as e:
        return f"[解析失败: {e}]"
    return ""


def _kb_dict(kb: KnowledgeBase, files: list[KnowledgeFile] | None = None) -> dict:
    d = {
        "id": kb.id, "name": kb.name, "description": kb.description or "",
        "tags": kb.tags or [], "status": kb.status,
        "created_at": kb.created_at.isoformat() if kb.created_at else "",
        "updated_at": kb.updated_at.isoformat() if kb.updated_at else "",
        "file_count": 0,
    }
    if files is not None:
        d["files"] = [_file_dict(f) for f in files]
        d["file_count"] = len(files)
    return d


def _file_dict(f: KnowledgeFile) -> dict:
    return {
        "id": f.id, "kb_id": f.kb_id, "name": f.name,
        "file_type": f.file_type, "size": f.size, "status": f.status,
        "created_at": f.created_at.isoformat() if f.created_at else "",
        "has_content": bool(f.parsed_content),
    }


# ── 知识库 CRUD ──────────────────────────────────────────────

@router.get("")
def list_kbs(q: str = Query(default=""), db: Session = Depends(get_db)):
    query = db.query(KnowledgeBase)
    if q:
        query = query.filter(KnowledgeBase.name.ilike(f"%{q}%"))
    kbs = query.order_by(KnowledgeBase.updated_at.desc()).all()
    result = []
    for kb in kbs:
        d = _kb_dict(kb)
        d["file_count"] = db.query(KnowledgeFile).filter(KnowledgeFile.kb_id == kb.id).count()
        result.append(d)
    return result


@router.post("", status_code=201)
def create_kb(name: str = Form(...), description: str = Form(""), tags: str = Form(""), db: Session = Depends(get_db)):
    tag_list = [t.strip() for t in tags.split(",") if t.strip()] if tags else []
    kb = KnowledgeBase(name=name, description=description, tags=tag_list)
    db.add(kb)
    db.commit()
    db.refresh(kb)
    return _kb_dict(kb, [])


@router.get("/search")
def search_kbs(q: str = Query(...), db: Session = Depends(get_db)):
    if not q.strip():
        return []
    files = db.query(KnowledgeFile).filter(KnowledgeFile.parsed_content.ilike(f"%{q}%")).limit(50).all()
    results = []
    seen_kb = {}
    for f in files:
        if f.kb_id not in seen_kb:
            kb = db.get(KnowledgeBase, f.kb_id)
            seen_kb[f.kb_id] = kb.name if kb else f.kb_id
        idx = f.parsed_content.lower().find(q.lower()) if f.parsed_content else -1
        snippet = f.parsed_content[max(0, idx-60):idx+120] if idx >= 0 and f.parsed_content else ""
        results.append({"kb_id": f.kb_id, "kb_name": seen_kb[f.kb_id], "file_id": f.id, "file_name": f.name, "snippet": snippet})
    return results


@router.get("/{kb_id}")
def get_kb(kb_id: str, db: Session = Depends(get_db)):
    kb = db.get(KnowledgeBase, kb_id)
    if not kb:
        raise HTTPException(404, "知识库不存在")
    files = db.query(KnowledgeFile).filter(KnowledgeFile.kb_id == kb_id).order_by(KnowledgeFile.created_at.desc()).all()
    return _kb_dict(kb, files)


@router.put("/{kb_id}")
def update_kb(kb_id: str, name: str = Form(None), description: str = Form(None), tags: str = Form(None), db: Session = Depends(get_db)):
    kb = db.get(KnowledgeBase, kb_id)
    if not kb:
        raise HTTPException(404, "知识库不存在")
    if name is not None:
        kb.name = name
    if description is not None:
        kb.description = description
    if tags is not None:
        kb.tags = [t.strip() for t in tags.split(",") if t.strip()]
    kb.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(kb)
    files = db.query(KnowledgeFile).filter(KnowledgeFile.kb_id == kb_id).all()
    return _kb_dict(kb, files)


@router.delete("/{kb_id}")
def delete_kb(kb_id: str, db: Session = Depends(get_db)):
    kb = db.get(KnowledgeBase, kb_id)
    if not kb:
        raise HTTPException(404, "知识库不存在")
    db.query(KnowledgeFile).filter(KnowledgeFile.kb_id == kb_id).delete()
    db.delete(kb)
    db.commit()
    return {"ok": True}


# ── 文件操作 ──────────────────────────────────────────────────

@router.post("/{kb_id}/files", status_code=201)
async def upload_file(kb_id: str, file: UploadFile = File(...), db: Session = Depends(get_db)):
    kb = db.get(KnowledgeBase, kb_id)
    if not kb:
        raise HTTPException(404, "知识库不存在")

    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    file_type = EXT_TYPE_MAP.get(ext, "other")
    save_name = f"{_uuid.uuid4().hex}.{ext}"
    save_path = os.path.join(UPLOAD_DIR, save_name)

    content = await file.read()
    with open(save_path, "wb") as fp:
        fp.write(content)

    parsed = ""
    if file_type in ("pdf", "word", "excel"):
        parsed = _parse_file(save_path, file_type)

    kf = KnowledgeFile(
        kb_id=kb_id,
        name=file.filename or save_name,
        file_type=file_type,
        file_path=save_path,
        size=len(content),
        parsed_content=parsed[:100000] if parsed else None,
        status="ready",
    )
    db.add(kf)
    kb.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(kf)
    return _file_dict(kf)


@router.delete("/{kb_id}/files/{fid}")
def delete_file(kb_id: str, fid: str, db: Session = Depends(get_db)):
    kf = db.get(KnowledgeFile, fid)
    if not kf or kf.kb_id != kb_id:
        raise HTTPException(404, "文件不存在")
    if kf.file_path and os.path.exists(kf.file_path):
        os.remove(kf.file_path)
    db.delete(kf)
    db.commit()
    return {"ok": True}


@router.get("/{kb_id}/files/{fid}/content")
def get_file_content(kb_id: str, fid: str, db: Session = Depends(get_db)):
    kf = db.get(KnowledgeFile, fid)
    if not kf or kf.kb_id != kb_id:
        raise HTTPException(404, "文件不存在")
    return {"content": kf.parsed_content or "", "file_type": kf.file_type, "name": kf.name}
