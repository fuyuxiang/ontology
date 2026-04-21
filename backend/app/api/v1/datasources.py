import socket
import logging
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.datasource import DataSource
from app.schemas.datasource import (
    DataSourceCreate, DataSourceUpdate, DataSourceDetail,
    DataSourceListItem, TestConnectionResult,
    TableListResult, TablePreviewResult, TableSchemaResult,
)
from app.services.datasource_utils import (
    get_connection as _get_connection,
    list_tables as _list_tables,
    preview_table as _preview_table,
    get_table_schema as _get_table_schema,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/datasources", tags=["datasources"])


@router.get("", response_model=list[DataSourceListItem])
def list_datasources(
    type: str | None = None,
    status: str | None = None,
    q: str | None = None,
    db: Session = Depends(get_db),
):
    query = db.query(DataSource)
    if type:
        query = query.filter(DataSource.type == type)
    if status:
        query = query.filter(DataSource.status == status)
    if q:
        query = query.filter(DataSource.name.ilike(f"%{q}%"))
    return query.order_by(DataSource.created_at.desc()).all()


@router.get("/{ds_id}", response_model=DataSourceDetail)
def get_datasource(ds_id: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    return ds


@router.post("", response_model=list[DataSourceDetail], status_code=201)
def create_datasource(body: DataSourceCreate, db: Session = Depends(get_db)):
    """连接数据库，获取所有表，为每张表创建一条数据源记录"""
    # 先用连接信息获取表列表
    tmp_ds = DataSource(**body.model_dump(), name="_tmp")
    try:
        tables = _list_tables(tmp_ds)
    except Exception as e:
        raise HTTPException(400, f"连接数据库失败: {e}")
    if not tables:
        raise HTTPException(400, "未获取到任何表，请检查连接信息和权限")

    created = []
    for tbl in tables:
        # 跳过已存在的同名记录
        if db.query(DataSource).filter(DataSource.name == tbl).first():
            continue
        ds = DataSource(
            **body.model_dump(),
            name=tbl,
            table_name=tbl,
            enabled=True,
        )
        db.add(ds)
        db.commit()
        db.refresh(ds)
        created.append(ds)
    return created


@router.put("/{ds_id}", response_model=DataSourceDetail)
def update_datasource(ds_id: str, body: DataSourceUpdate, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(ds, k, v)
    db.commit()
    db.refresh(ds)
    return ds


@router.delete("/{ds_id}", status_code=204)
def delete_datasource(ds_id: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    db.delete(ds)
    db.commit()


@router.post("/{ds_id}/test", response_model=TestConnectionResult)
def test_connection(ds_id: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    return _test_conn(ds.host, ds.port)


@router.post("/test", response_model=TestConnectionResult)
def test_connection_inline(body: DataSourceCreate):
    return _test_conn(body.host, body.port)


@router.post("/fetch-tables", response_model=TableListResult)
def fetch_tables_inline(body: DataSourceCreate):
    """根据连接信息动态获取数据库中的表列表（无需先创建数据源）"""
    ds = DataSource(**body.model_dump(), name="_tmp")
    try:
        tables = _list_tables(ds)
        return {"tables": tables}
    except Exception as e:
        raise HTTPException(400, f"获取表列表失败: {e}")


def _test_conn(host: str, port: int) -> dict:
    try:
        sock = socket.create_connection((host, port), timeout=5)
        sock.close()
        return {"success": True, "message": "连接成功"}
    except Exception as e:
        return {"success": False, "message": f"连接失败: {e}"}


@router.post("/{ds_id}/toggle", response_model=DataSourceDetail)
def toggle_datasource(ds_id: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    ds.enabled = not ds.enabled
    db.commit()
    db.refresh(ds)
    return ds


@router.post("/{ds_id}/refresh-tables", response_model=DataSourceDetail)
def refresh_tables(ds_id: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    if not ds.enabled:
        raise HTTPException(403, "数据源管道已关闭，无法同步数据。请在数据工坊->数据源管理中启用该管道")
    if not ds.table_name:
        raise HTTPException(400, "该数据源未关联数据表，请先编辑选择数据表")
    count, err = _count_records(ds)
    if err:
        raise HTTPException(400, f"同步失败: {err}")
    ds.record_count = count
    db.commit()
    db.refresh(ds)
    return ds


def _count_records(ds: DataSource) -> tuple[int, str | None]:
    """连接数据库并查询当前数据表的记录条数，返回 (count, error_msg)"""
    try:
        conn = _get_connection(ds)
        if not conn:
            return 0, f"不支持的数据源类型: {ds.type}"
        cur = conn.cursor()
        quote = "`" if ds.type == "mysql" else '"'
        cur.execute(f"SELECT COUNT(*) FROM {quote}{ds.table_name}{quote}")
        count = cur.fetchone()[0]
        cur.close()
        conn.close()
        return count, None
    except Exception as e:
        logger.warning(f"获取记录条数失败 [{ds.name}]: {e}")
        return 0, str(e)


@router.get("/{ds_id}/tables", response_model=TableListResult)
def get_tables(ds_id: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    try:
        tables = _list_tables(ds)
        return {"tables": tables}
    except Exception as e:
        raise HTTPException(400, f"获取表列表失败: {e}")


@router.get("/{ds_id}/preview", response_model=TablePreviewResult)
def preview_datasource(ds_id: str, db: Session = Depends(get_db)):
    """直接预览该数据源关联表的前20条数据"""
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    if not ds.enabled:
        raise HTTPException(403, "数据源管道已关闭，无法读取数据。请在数据工坊->数据源管理中启用该管道")
    if not ds.table_name:
        raise HTTPException(400, "该数据源未关联数据表")
    try:
        return _preview_table(ds, ds.table_name)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"查询失败: {e}")


@router.get("/{ds_id}/tables/{table_name}/preview", response_model=TablePreviewResult)
def preview_table(ds_id: str, table_name: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    if not ds.enabled:
        raise HTTPException(403, "数据源管道已关闭，无法读取数据。请在数据工坊->数据源管理中启用该管道")
    try:
        return _preview_table(ds, table_name)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"查询失败: {e}")


@router.get("/{ds_id}/tables/{table_name}/schema", response_model=TableSchemaResult)
def get_table_schema(ds_id: str, table_name: str, db: Session = Depends(get_db)):
    ds = db.get(DataSource, ds_id)
    if not ds:
        raise HTTPException(404, "数据源不存在")
    try:
        columns = _get_table_schema(ds, table_name)
        return {"table": table_name, "columns": columns}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"获取表结构失败: {e}")


# ── 多模态数据源 ──────────────────────────────────────────────

import os, uuid as _uuid
from fastapi import UploadFile, File, Form
from pydantic import BaseModel as _BM
from typing import Any as _Any

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "..", "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


def _parse_file(path: str, file_type: str) -> str:
    try:
        if file_type == "pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif file_type == "word":
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif file_type == "excel":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)
    except Exception as e:
        return f"[解析失败: {e}]"
    return ""


@router.post("/upload", status_code=201)
async def upload_file_source(
    file: UploadFile = File(...),
    name: str = Form(...),
    description: str = Form(""),
    db: Session = Depends(get_db),
):
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    type_map = {"pdf": "pdf", "doc": "word", "docx": "word",
                "xls": "excel", "xlsx": "excel",
                "png": "image", "jpg": "image", "jpeg": "image", "gif": "image",
                "mp4": "video", "avi": "video", "mov": "video"}
    file_type = type_map.get(ext, "other")
    ds_type = "file"

    save_name = f"{_uuid.uuid4().hex}.{ext}"
    save_path = os.path.join(UPLOAD_DIR, save_name)
    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)

    parsed = ""
    if file_type in ("pdf", "word", "excel"):
        parsed = _parse_file(save_path, file_type)

    ds = DataSource(
        name=name or file.filename,
        type=ds_type,
        host="local",
        port=0,
        source_category="file",
        file_path=save_path,
        file_type=file_type,
        parsed_content=parsed[:50000] if parsed else None,
        description=description,
        status="active",
        enabled=True,
        record_count=len(parsed.splitlines()) if parsed else 0,
    )
    db.add(ds)
    db.commit()
    db.refresh(ds)
    return {"id": ds.id, "name": ds.name, "file_type": file_type, "parsed_lines": ds.record_count}


class ApiSourceBody(_BM):
    name: str
    api_url: str
    api_method: str = "GET"
    api_headers: dict | None = None
    api_body: str | None = None
    poll_interval: int = 60
    description: str = ""


@router.post("/api-source", status_code=201)
def create_api_source(body: ApiSourceBody, db: Session = Depends(get_db)):
    import httpx
    try:
        resp = httpx.request(
            body.api_method, body.api_url,
            headers=body.api_headers or {},
            content=body.api_body,
            timeout=10,
        )
        status = "active" if resp.status_code < 400 else "error"
        preview = resp.text[:2000]
    except Exception as e:
        status = "error"
        preview = str(e)

    ds = DataSource(
        name=body.name,
        type="api",
        host=body.api_url.split("/")[2] if "//" in body.api_url else body.api_url,
        port=443,
        source_category="api",
        api_url=body.api_url,
        api_method=body.api_method,
        api_headers=body.api_headers,
        api_body=body.api_body,
        poll_interval=body.poll_interval,
        parsed_content=preview,
        description=body.description,
        status=status,
        enabled=True,
    )
    db.add(ds)
    db.commit()
    db.refresh(ds)
    return {"id": ds.id, "name": ds.name, "status": status, "preview": preview}


class MqSourceBody(_BM):
    name: str
    host: str
    port: int = 9092
    mq_topic: str
    mq_group: str = "ontology-consumer"
    username: str = ""
    password: str = ""
    poll_interval: int = 60
    description: str = ""


@router.post("/mq-source", status_code=201)
def create_mq_source(body: MqSourceBody, db: Session = Depends(get_db)):
    ds = DataSource(
        name=body.name,
        type="kafka",
        host=body.host,
        port=body.port,
        username=body.username,
        password=body.password,
        source_category="mq",
        mq_topic=body.mq_topic,
        mq_group=body.mq_group,
        poll_interval=body.poll_interval,
        description=body.description,
        status="active",
        enabled=True,
    )
    db.add(ds)
    db.commit()
    db.refresh(ds)
    return {"id": ds.id, "name": ds.name, "status": "active"}
